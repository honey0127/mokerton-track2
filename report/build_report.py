"""예선 결과 보고서 초안을 주최측 양식(form.docx)에 채운다.

수치는 전부 파일에서 읽는다:
  - 정확도: eval_val 채점 JSON (/home/claude/eval_val_cloud)
  - 보조 속도: i7-10510U 15라운드 (benchmarks/bench_key15.jsonl)
  - 주 속도: M1 Pro 결과가 있으면 M1DIR/bench_key15_m1pro.jsonl, 없으면 【M1】 자리표시
【...】 로 감싼 글자는 노란 형광으로 표시된다 (사람이 채울 곳).
"""
import copy, json, re, sys
from pathlib import Path
import docx
from docx.shared import Pt
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).parent
REPO = Path("/home/claude/mokerton-track2")
EV = Path("/home/claude/eval_val_cloud")
M1DIR = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/home/claude/m1_results")
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("/mnt/user-data/outputs/트랙2_예선결과보고서_초안.docx")

A, B, C, D = "best_1280x384", "best-p34_1280x384", "best_640x192", "best-p34_640x192"
I8 = "int8-mixed-percentile"


def ev(stem, var):
    j = json.loads((EV / f"ap40_eval_val_{stem}_op13_{var}.json").read_text())
    return {"Car": j["Car/Moderate"]["AP40"], "Pedestrian": j["Pedestrian/Moderate"]["AP40"],
            "Cyclist": j["Cyclist/Moderate"]["AP40"], "mAP": j["mAP/Moderate"]}


def hold(stem, var):
    return json.loads((REPO / f"benchmarks/ap40_holdout_{stem}_op13_{var}.json").read_text())["mAP/Moderate"]


def load_bench(path):
    out = {}
    if path.exists():
        for ln in path.read_text().splitlines():
            if ln.strip():
                r = json.loads(ln)
                if r.get("kind") == "bench" and int(r["threads"]) == 4:
                    out[Path(r["model"]).stem] = r
    return out


LAP = load_bench(REPO / "benchmarks/bench_key15.jsonl")
M1 = load_bench(M1DIR / "bench_key15_m1pro.jsonl")
M1ENV = {}
if (M1DIR / "env_m1.txt").exists():
    for ln in (M1DIR / "env_m1.txt").read_text().splitlines():
        if ":" in ln:
            k, v = ln.split(":", 1)
            M1ENV[k.strip()] = v.strip()
HAVE_M1 = len(M1) >= 8

SIZE = {k: round((REPO / f"models/{k}.onnx").stat().st_size / 1e6, 2) for k in
        [f"{s}_op13_{v}" for s in (A, B, C, D) for v in ("fp32", I8)]}
PARAMS = {A: 2.583, B: 1.845, C: 2.583, D: 1.845}
GFLOPS = {A: 7.66, B: 6.96, C: 1.90, D: 1.72}


def f2(x):
    return f"{x:.2f}"


def pct(a, b):  # a -> b 변화율
    return f"{(b - a) / a * 100:+.1f}%"


def m1(stem, var, key):
    r = M1.get(f"{stem}_op13_{var}")
    if not r:
        return None
    return {"p50": r["p50_ms"], "p99": r["p99_ms"], "fps": round(1000 / r["p50_ms"], 1), "rss": r["peak_rss_MB"]}[key]


def m1r(s0, v0, s1, v1):
    return m1(s0, v0, "p50") / m1(s1, v1, "p50")


def m1s(stem, var, key, fmt="{:.2f}"):
    v = m1(stem, var, key)
    return fmt.format(v) if v is not None else "【M1】"


# ---------------------------------------------------------------- docx helpers
d = docx.Document(HERE / "form.docx")
body = d.element.body
HL = re.compile(r"(【[^】]*】)")


def add_runs(p, text, size=None, bold=False, mono=False, italic=False):
    for part in HL.split(text):
        if not part:
            continue
        r = p.add_run(part)
        if size:
            r.font.size = Pt(size)
        r.bold = bold or None
        r.italic = italic or None
        if mono:
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        if part.startswith("【"):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def set_cell(cell, text, size=9, bold=False, align=None):
    ps = cell.paragraphs
    p = ps[0]
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    lines = text.split("\n")
    add_runs(p, lines[0], size, bold)
    if align is not None:
        p.alignment = align
    for ln in lines[1:]:
        np_ = copy.deepcopy(p._element)
        for r in np_.findall(qn("w:r")):
            np_.remove(r)
        p._element.addnext(np_)
        p = docx.text.paragraph.Paragraph(np_, cell)
        add_runs(p, ln, size, bold)


def fill_box(table, items, size=9.5):
    """1x1 서술 상자. items: str | ('h', 굵은 소제목) | ('b', 굵은 머리, 나머지) | ('code', 여러 줄) | ('li', 글머리 문장)"""
    cell = table.rows[0].cells[0]
    ps = cell.paragraphs
    proto = copy.deepcopy(ps[0]._element)
    for r in proto.findall(qn("w:r")):
        proto.remove(r)
    for p in ps:
        p._element.getparent().remove(p._element)
    tc = cell._tc

    def newp():
        e = copy.deepcopy(proto)
        tc.append(e)
        return docx.text.paragraph.Paragraph(e, cell)

    for it in items:
        if isinstance(it, str):
            add_runs(newp(), it, size)
        elif it[0] == "h":
            add_runs(newp(), it[1], size, bold=True)
        elif it[0] == "b":
            p = newp()
            add_runs(p, it[1], size, bold=True)
            add_runs(p, it[2], size)
        elif it[0] == "li":
            p = newp()
            p.paragraph_format.left_indent = Pt(10)
            p.paragraph_format.first_line_indent = Pt(-8)
            add_runs(p, "· " + it[1], size)
        elif it[0] == "code":
            for ln in it[1].strip("\n").split("\n"):
                p = newp()
                p.paragraph_format.line_spacing = 1.0
                add_runs(p, ln if ln else " ", 8, mono=True)


def elems():
    return list(body.iterchildren())


def tables_in_order():
    return [docx.table.Table(e, d) for e in elems() if e.tag == qn("w:tbl")]


T = tables_in_order()
# 인덱스: 0 안내, 1 팀정보, 2 결과요약, 3 요약상자, 4 대상모델, 5 선정사유, 6 Hailo, 7 환경, 8 기법요약,
#        9 양자화상세, 10 근거상자, 11 벤치규격, 12 프로토콜, 13 정확도, 14 속도, 15 해석, 16 한계, 17 이식계획,
#        18 체크리스트, 19 환경구성, 20 실행명령, 21 기대출력, 22 유의사항
assert len(T) == 23, len(T)


def ensure_rows(t, n):
    while len(t.rows) < n:
        t._tbl.append(copy.deepcopy(t.rows[-1]._tr))


def clone_after(anchor_el, t, caption=None):
    """표 t 를 복제해 anchor 뒤에 넣는다. caption 이 있으면 그 앞에 굵은 문단을 둔다."""
    new = copy.deepcopy(t._tbl)
    spacer = OxmlElement("w:p")
    anchor_el.addnext(spacer)
    last = spacer
    if caption:
        cap = OxmlElement("w:p")
        last.addnext(cap)
        cp = docx.text.paragraph.Paragraph(cap, d)
        add_runs(cp, caption, 9.5, bold=True)
        cp.paragraph_format.keep_with_next = True
        last = cap
    last.addnext(new)
    return docx.table.Table(new, d)


def caption_before(t, text):
    p = OxmlElement("w:p")
    t._tbl.addprevious(p)
    pp = docx.text.paragraph.Paragraph(p, d)
    add_runs(pp, text, 9.5, bold=True)
    pp.paragraph_format.keep_with_next = True


def fill_grid(t, rows, header=None, center_from=1):
    ensure_rows(t, len(rows) + 1)
    while len(t.rows) > len(rows) + 1:          # 복제한 표의 남는 행 삭제
        t._tbl.remove(t.rows[-1]._tr)
    if header:
        for j, h in enumerate(header):
            set_cell(t.rows[0].cells[j], h, bold=True, align=1)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            set_cell(t.rows[i].cells[j], v, align=1 if j >= center_from else None)


# 회색 기울임 안내문 삭제 (양식 지시)
for p in list(d.paragraphs):
    if p.text.strip().startswith("▸") or p.text.strip().startswith("요약 : 어떤 기법"):
        p._element.getparent().remove(p._element)

# 안내 상자(작성 안내) 삭제
T[0]._tbl.getparent().remove(T[0]._tbl)

# ================================================================ 데이터
eA, eB, eC, eD = ev(A, "fp32"), ev(B, I8), ev(C, "fp32"), ev(D, I8)
eBf, eDf, eAi, eCi = ev(B, "fp32"), ev(D, "fp32"), ev(A, I8), ev(C, I8)

# ================================================================ 1. 팀 정보
t = T[1]
set_cell(t.rows[1].cells[1], "【팀명】")
set_cell(t.rows[3].cells[1], "YOLO11n (Ultralytics 8.4.135) — KITTI 3클래스 fine-tuning, P5 검출 헤드 제거(P3·P4 2-scale) + INT8 혼합 정밀도 PTQ\n"
         "정밀도 프로파일 1280×384 = 모델 B / 초경량 프로파일 640×192 = 모델 D")
set_cell(t.rows[4].cells[1], "【성명 / 학번 / 역할】")
set_cell(t.rows[5].cells[1], "【성명 / 학번 / 역할】")
set_cell(t.rows[6].cells[1], "【성명 / 학번 / 역할】")
set_cell(t.rows[7].cells[1], "【이메일 / 전화】")
set_cell(t.rows[8].cells[1], "2026년 【 】월 【 】일")

# ================================================================ 2. 결과 요약 (A FP32 -> B INT8)
t = T[2]
lat = lambda s, v: m1(s, v, "p50")
rows2 = [
    ("모델 크기 (MB)", f2(SIZE[f'{A}_op13_fp32']), f2(SIZE[f'{B}_op13_{I8}']), pct(SIZE[f'{A}_op13_fp32'], SIZE[f'{B}_op13_{I8}'])),
    ("파라미터 수 (M)", f"{PARAMS[A]:.3f}", f"{PARAMS[B]:.3f}", pct(PARAMS[A], PARAMS[B])),
    ("연산량 (GFLOPs)", f2(GFLOPS[A]), f2(GFLOPS[B]), pct(GFLOPS[A], GFLOPS[B])),
]
if HAVE_M1:
    rows2 += [("추론 Latency (ms)", f2(m1(A, 'fp32', 'p50')), f2(m1(B, I8, 'p50')), pct(m1(A, 'fp32', 'p50'), m1(B, I8, 'p50'))),
              ("Throughput (FPS)", f"{m1(A, 'fp32', 'fps'):.1f}", f"{m1(B, I8, 'fps'):.1f}", pct(m1(A, 'fp32', 'fps'), m1(B, I8, 'fps'))),
              ("피크 메모리 사용량 (MB)", f"{m1(A, 'fp32', 'rss'):.1f}", f"{m1(B, I8, 'rss'):.1f}", pct(m1(A, 'fp32', 'rss'), m1(B, I8, 'rss')))]
else:
    rows2 += [("추론 Latency (ms)", "【M1】", "【M1】", "【M1】"), ("Throughput (FPS)", "【M1】", "【M1】", "【M1】"),
              ("피크 메모리 사용량 (MB)", "【M1】", "【M1】", "【M1】")]
rows2.append(("정확도", f"{eA['mAP']:.2f}", f"{eB['mAP']:.2f}", f"{eB['mAP'] - eA['mAP']:+.2f}%p"))
for i, r in enumerate(rows2, start=1):
    for j in (1, 2, 3):
        set_cell(t.rows[i].cells[j], r[j], align=1)
set_cell(t.rows[7].cells[0], "정확도 (KITTI Moderate AP40 mAP, eval_val 1,000장)")
caption_before(t, "정밀도 프로파일 1280×384 — Baseline: A (YOLO11n FP32) / 최적화 후: B (P5 제거 + INT8)")

lat_txt = (f"지연(M1 Pro, 4스레드 p50)은 {m1(A, 'fp32', 'p50'):.2f} → {m1(B, I8, 'p50'):.2f}ms ({pct(m1(A, 'fp32', 'p50'), m1(B, I8, 'p50'))})"
           if HAVE_M1 else "지연(M1 Pro, 4스레드 p50)은 【M1】ms → 【M1】ms")
lat_txt_cd = (f"{m1(C, 'fp32', 'p50'):.2f} → {m1(D, I8, 'p50'):.2f}ms" if HAVE_M1 else "【M1】ms → 【M1】ms")
fill_box(T[3], [
    ("b", "요약 : ", f"KITTI로 fine-tuning한 YOLO11n FP32 모델(A, 1280×384)을 기준으로 두 가지를 적용했다. "
     f"① P5(stride 32) 검출 헤드를 떼고 다시 학습해 파라미터를 {PARAMS[A]:.3f}M → {PARAMS[B]:.3f}M({pct(PARAMS[A], PARAMS[B])}), "
     f"연산량을 {GFLOPS[A]:.2f} → {GFLOPS[B]:.2f} GFLOPs({pct(GFLOPS[A], GFLOPS[B])})로 줄였다. "
     f"② Detect 디코드 경로만 FP32로 남긴 INT8 정적 PTQ(Percentile 캘리브레이션, 가중치 7비트)를 적용했다."),
    f"그 결과 모델 크기는 {SIZE[f'{A}_op13_fp32']:.2f} → {SIZE[f'{B}_op13_{I8}']:.2f}MB({pct(SIZE[f'{A}_op13_fp32'], SIZE[f'{B}_op13_{I8}'])})로 줄었고, {lat_txt}다. "
    f"공식 평가셋(eval_val 1,000장) Moderate AP40 mAP는 {eA['mAP']:.2f} → {eB['mAP']:.2f}({eB['mAP'] - eA['mAP']:+.2f}%p)다.",
    f"초경량 프로파일(640×192, C → D)에도 같은 방법을 적용했다. 크기 {SIZE[f'{C}_op13_fp32']:.2f} → {SIZE[f'{D}_op13_{I8}']:.2f}MB, "
    f"연산량 {GFLOPS[C]:.2f} → {GFLOPS[D]:.2f} GFLOPs, 지연 {lat_txt_cd}, mAP {eC['mAP']:.2f} → {eD['mAP']:.2f}({eD['mAP'] - eC['mAP']:+.2f}%p).",
    "주력 모델(B, D)과 INT8 방식은 eval_val을 한 번도 채점하지 않은 상태에서 holdout 500장으로 정했다. 결정 규칙과 결정 내용은 저장소 DECISION.md(커밋 fcba017)에 먼저 커밋했다. "
    "eval_val은 22개 모델 전체를 한 번에, 1회만 채점했다. 이후 2차 라운드(채널 프루닝, 양자화 설정 개선)를 같은 방식으로 규칙을 먼저 정해 시도했으나 "
    "채택 조건을 통과한 후보가 없어(5.3 ⑨, DECISION_R2.md) 위 수치가 최종이다.",
])

# ================================================================ 3.1 대상 모델
t = T[4]
set_cell(t.rows[1].cells[1], "YOLO11n (Ultralytics 8.4.135, scale n) — KITTI 3클래스(Car / Pedestrian / Cyclist)\n"
         "변형: P5 검출 헤드 제거 'yolo11n-p34' (cfg/yolo11n-p34.yaml, scripts/make_p34.py)")
set_cell(t.rows[2].cells[1], "https://github.com/honey0127/mokerton-track2\n(기반 모델: https://github.com/ultralytics/ultralytics)")
set_cell(t.rows[3].cells[1], "b0ec2cd — 모델 확정 및 eval_val 채점 직전 커밋 (ONNX 22개 + 학습 가중치 SHA-256 고정: MODELS.sha256)\n최종 제출 커밋: 【커밋 해시】")
set_cell(t.rows[4].cells[1], "Ultralytics 공식 yolo11n.pt (COCO 사전학습). A·C는 그대로, B·D는 P5 부분을 뺀 같은 가중치(층 0~19 + Detect P3·P4 가지)에서 출발해 "
         "KITTI train 5,981장으로 fine-tuning (eval_val·holdout 제외)")
set_cell(t.rows[5].cells[1], "KITTI 원본 1242×375 (일부 1224×370 등) → 모델 입력 1280×384 (정밀도) / 640×192 (초경량), 비율 유지 letterbox")

# ================================================================ 3.2 선정 사유
fill_box(T[5], [
    ("b", "엣지 대상에 맞는 크기. ", "본선 보드는 Raspberry Pi 5(4코어) + Hailo-10H다. YOLO11n은 2.6M 파라미터의 nano급 모델로, "
     "YOLO 계열에서 가장 작은 축이면서 COCO 사전학습 가중치가 공개되어 있어 6천 장 남짓한 KITTI에서도 fine-tuning만으로 높은 정확도를 낸다 "
     f"(eval_val Moderate mAP: 1280×384 FP32 {eA['mAP']:.2f})."),
    ("b", "Hailo DFC 호환성이 확인된 구조. ", "Hailo Model Zoo가 YOLO11n을 Hailo-10H 대상으로 제공한다(참고치 303 FPS). "
     "우리도 COCO 가중치의 YOLO11n 640×640을 DFC 5.x로 hailo10h HEF까지 컴파일해 통과를 확인했다(hailo/compiled_640.html, 4 context)."),
    ("b", "후처리를 떼어내기 쉬운 출력 구조. ", "앵커 없는 decoupled head라 출력이 [1, 4+3, N] 텐서 하나다. 박스 디코드(DFL)·NMS를 그래프 밖으로 빼는 지점이 "
     "명확하다. 이 지점이 ORT INT8의 FP32 유지 경계(5장)이자 Hailo parse의 절단 지점(3.3)이 된다."),
    ("b", "해상도. ", "KITTI 원본은 1242×375로 가로가 3.3배 길다. 정사각 640×640은 letterbox 패딩에 연산의 약 70%를 낭비한다. 공지 권장(가로 ≤1280, 32의 배수)대로 "
     "1280×384(패딩 거의 0, 1242×375 → 1272×384)와 그 1/4 연산의 640×192를 프로파일로 정했다."),
    ("b", "P5 헤드 제거를 시도한 이유. ", "train+holdout 6,481장의 Moderate GT 17,671개를 크기별로 세어 보았다(results/box_sizes.md, eval_val 미사용). "
     "P5(stride 32)가 맡는 큰 물체는 640×192에서 1.28%, 1280×384에서도 16.69%뿐이었다. "
     "작은 물체 위주인 데이터에서 P5 경로는 연산 대비 기여가 작다고 보고 떼어 낸 뒤 다시 학습했다."),
])

# ================================================================ 3.3 Hailo
fill_box(T[6], [
    ("b", "구성 연산. ", "Conv·BN(fused)·SiLU·Concat·Upsample(nearest)·MaxPool(SPPF)·Split, 그리고 백본 끝 C2PSA 블록의 attention(MatMul·Softmax)이다. "
     "C2PSA는 B·D에도 남는다(백본 P5 단계는 유지). 같은 블록을 포함한 YOLO11n 640×640이 DFC 5.x에서 컴파일되었으므로 비표준 레이어는 없다."),
    ("b", "디코드 경로는 칩 밖으로. ", "DFL softmax·앵커 생성·박스 디코드·NMS는 NPU에 넣지 않는다. hailo parser의 end node를 Detect의 마지막 Conv(cv2.i.2 박스, cv3.i.2 클래스)로 지정해 "
     "그래프를 자르고, 후처리는 CPU에서 한다. Ultralytics 8.4.135의 Hailo export(engine/exporter.py export_hailo)도 같은 지점에서 자른다."),
    ("b", "P5 제거의 영향. ", "end node가 6개(3-scale)에서 4개(P3·P4: /model.20/cv2.0.2, cv3.0.2, cv2.1.2, cv3.1.2)로 바뀌어 Model Zoo 설정을 그대로 쓸 수 없다. "
     "hailo/01_export_onnx.py가 모델마다 end node를 자동으로 찾아 기록한다."),
    ("b", "입력. ", "batch 1, 정적 형상, opset 13. 직사각 입력(384×1280)은 DFC가 지원하며, 정규화는 .alls의 normalization 층으로 NPU 안에서 한다. "
     "보정 이미지는 train에서만 뽑고, eval_val이 섞이면 스크립트가 중단된다(hailo/02_make_calib.py)."),
    ("b", "남은 위험과 대응. ", "① 1280×384는 640×640보다 픽셀이 1.2배라 context 수·리소스가 늘 수 있다. 리소스 초과 시 640×192(D)를 우선 이식한다. "
     "② DFC는 x86 Linux 전용이고 WSL2에서는 GPU를 쓰지 못해 고급 최적화(level 3~4, AdaRound)가 느리다. 학과 공용 PC(RTX 5070)의 네이티브 Linux나 Colab을 쓴다. "
     "③ B·D의 HEF 컴파일은 아직 하지 않았다(7.2 계획 1주차)."),
])

# ================================================================ 4. 환경
t = T[7]
m1cpu = M1ENV.get("cpu", "Apple M1 Pro")
m1cores = M1ENV.get("cores", "【P/E 코어 수】").replace("P=", "P코어 ").replace(" E=", " + E코어 ").split(" logical")[0]
set_cell(t.rows[1].cells[1], f"[속도·메모리 측정] {m1cpu} ({m1cores}), arm64\n"
         "[정확도 채점] Intel Xeon 2.80GHz, AVX-512 VNNI, 2 vCPU (클라우드 Linux 컨테이너)\n"
         "[보조 측정] Intel Core i7-10510U, 4코어 8스레드, AVX2 (VNNI 없음)\n"
         "[학습] Intel Core i9-14900KF")
set_cell(t.rows[2].cells[1], "[학습] NVIDIA GeForce RTX 5070 12GB (driver 591.86)\n[측정·채점] 사용 안 함 — 모든 추론은 ONNX Runtime CPUExecutionProvider")
set_cell(t.rows[3].cells[1], f"[M1] {M1ENV.get('memory_GB', '【 】')}GB  /  [채점] 8GB  /  [보조 i7] 8GB  /  [학습] 32GB (31.8GB 표시)")
set_cell(t.rows[4].cells[1], f"[M1] macOS {M1ENV.get('macos', '【버전】')}  /  [채점] Ubuntu 24.04.5 LTS  /  [보조 i7] Windows 11 Pro 10.0.26100  /  [학습] Windows 11")
set_cell(t.rows[5].cells[1], f"[M1] {M1ENV.get('python', '【 】')} (uv 설치 CPython)  /  [채점] 3.13.16  /  [보조 i7] 3.12.10")
set_cell(t.rows[6].cells[1], "[학습] PyTorch 2.11.0+cu128, Ultralytics (best.pt 생성)\n"
         "[내보내기] Ultralytics 8.4.135, PyTorch 2.13.0+cpu, ONNX 1.22.0 (opset 13)\n"
         f"[추론·측정·채점] ONNX Runtime 1.29.0 (CPUExecutionProvider), numpy {M1ENV.get('numpy', '')}{' (M1)' if M1ENV.get('numpy') else ''} / 2.5.4 (채점), numba 0.68.0 (AP40 판정 엔진)")
set_cell(t.rows[7].cells[1], "[학습] CUDA 12.8 (PyTorch cu128 휠 내장), cuDNN 【버전】\n[측정·채점] CUDA 미사용")
set_cell(t.rows[8].cells[1], "ONNX Runtime 1.29.0 quantization (quantize_static, QDQ) — scripts/quantize_ort.py\n"
         "본선용: Hailo Dataflow Compiler 5.x (hailo10h)")

# ================================================================ 5.1 기법 요약
t = T[8]
ensure_rows(t, 7)
h1 = hold(A, "fp32")
r51 = [
    ("P5 검출 헤드 제거 + 재학습", "Neck P5 출력(층 20~22)과 Detect P5 가지. 백본은 유지",
     "Detect 3-scale → 2-scale(P3·P4). A·C와 같은 COCO 가중치·학습 설정",
     f"Params {pct(PARAMS[A], PARAMS[B])}, GFLOPs {pct(GFLOPS[A], GFLOPS[B])}(1280) / {pct(GFLOPS[C], GFLOPS[D])}(640). "
     f"FP32 mAP {eA['mAP']:.2f}→{eBf['mAP']:.2f}(1280), {eC['mAP']:.2f}→{eDf['mAP']:.2f}(640)"),
    ("INT8 정적 PTQ", "Conv 등 그래프 전체 (아래 제외 경로 빼고)",
     "QDQ, 활성 uint8 비대칭 / 가중치 int8 대칭 per-channel, 보정 256장",
     f"크기 {SIZE[f'{B}_op13_fp32']:.2f}→{SIZE[f'{B}_op13_{I8}']:.2f}MB({pct(SIZE[f'{B}_op13_fp32'], SIZE[f'{B}_op13_{I8}'])}), B 기준"),
    ("혼합 정밀도: Detect 디코드 경로 FP32", "Detect 내부 non-Conv 연산 + 출력 생성 노드 (B·D 45개 / A·C 59개)",
     "scripts/head_nodes.py 자동 추출 → nodes_to_exclude",
     f"전체 INT8은 mAP 0.00 → 혼합 {ev(B, 'int8-mixed')['mAP']:.2f} (B, MinMax)"),
    ("Percentile 캘리브레이션", "활성값 스케일", "ORT Percentile(99.999), 8장 단위 분할 수집",
     f"B: MinMax {ev(B, 'int8-mixed')['mAP']:.2f} → {eB['mAP']:.2f} / D: {ev(D, 'int8-mixed')['mAP']:.2f} → {eD['mAP']:.2f} (eval_val)"),
    ("가중치 7비트 (reduce_range)", "양자화된 Conv 가중치", "가중치 범위 [-127,127] → [-64,64]",
     "VNNI 없는 AVX2 CPU의 int16 포화 제거. CPU 간 INT8 점수 차 1~3점 → 0.3점 이내 (holdout)"),
    ("직사각 입력 해상도 프로파일", "입력", "1280×384 / 640×192, letterbox",
     f"640×192는 GFLOPs {GFLOPS[A]:.2f}→{GFLOPS[C]:.2f}(1/4), FP32 mAP {eA['mAP']:.2f}→{eC['mAP']:.2f}"),
]
for i, r in enumerate(r51, start=1):
    for j, v in enumerate(r):
        set_cell(t.rows[i].cells[j], v)

# ================================================================ 5.2 양자화 상세
t = T[9]
set_cell(t.rows[1].cells[1], "PTQ — 정적(static) 사후 양자화, ONNX Runtime quantize_static, QDQ 형식. QAT는 적용하지 않음")
set_cell(t.rows[2].cells[1], "혼합 — INT8(활성 uint8 비대칭 / 가중치 int8 대칭 per-channel, 7비트 범위 [-64,64]) + Detect 디코드 경로 FP32")
set_cell(t.rows[3].cells[1], "KITTI training 중 우리 train split(splits/train.txt)에서 무작위 추출. eval_val·holdout 제외. 추론과 같은 letterbox 전처리")
set_cell(t.rows[4].cells[1], "256장 (seed 0), 8장 단위로 끊어 수집(CalibStridedMinMax=8) — 히스토그램 보정의 메모리 폭증 방지")
set_cell(t.rows[5].cells[1], "Detect 모듈의 디코드 경로: DFL softmax·Reshape·Concat·앵커 연산 등 non-Conv 연산과 그래프 출력 생성 노드\n"
         "B·D 45개, A·C 59개 (models/exclude_*.txt). Detect의 마지막 Conv(cv2.*.2, cv3.*.2)와 DFL Conv는 INT8")

# ================================================================ 5.3 근거 및 과정
bMin, dMin = ev(B, "int8-mixed"), ev(D, "int8-mixed")
fill_box(T[10], [
    ("h", "① 해상도와 학습"),
    "KITTI 원본 비율(3.3:1)에 맞춰 1280×384와 640×192를 정했다. 학습은 정사각 캔버스 + mosaic 증강으로 했다. Ultralytics는 rect=True면 mosaic을 강제로 끄는데, "
    "6천 장 규모에서는 과적합 위험이 커서다. 합성곱 네트워크라 정사각으로 학습해도 직사각으로 추론할 수 있다. 공식 판정이 Van·Person_sitting을 '무시'로 처리하므로 "
    "학습 라벨에서 Van→Car, Person_sitting→Pedestrian으로 병합해 recall을 얻었다. 모델 선택은 holdout 500장으로만 했다.",
    ("h", "② 첫 시도: 전체 INT8 → mAP 0.00 (실패)"),
    "출력 텐서 [1, 7, N]은 0~3행이 박스 좌표(0~1280 픽셀), 4~6행이 클래스 확률(0~1)이다. 단일 스케일로 양자화하면 스케일이 약 1280/255로 잡혀 클래스 점수가 전부 0으로 뭉개진다. "
    "FP32와 출력을 비교하니 클래스 행의 고유값이 1822개에서 1개로 줄어 있었다(scripts/diag_outputs.py). 22개 모델 중 전체 INT8 4개는 eval_val에서도 모두 0.00이다.",
    ("h", "③ 두 번째 시도: Conv만 INT8 → 정확도는 지키지만 FP32보다 느림"),
    "Conv 사이 연산이 FP32로 남으면 Conv 경계마다 Quantize/Dequantize가 끼어 ORT가 양자화 서브그래프를 합치지 못한다. "
    "holdout 측정에서 FP32 대비 0.77~0.81배로 오히려 느렸고 peak RSS도 FP32보다 컸다.",
    ("h", "④ 채택: Detect 디코드 경로만 FP32"),
    "head_nodes.py가 Detect 접두사를 찾아 내부 non-Conv 연산과 출력 노드만 제외한다. 백본·넥·헤드 Conv는 모두 INT8이라 속도를 지키면서 박스·확률의 스케일 충돌을 피한다. "
    "이 경계는 Ultralytics의 Hailo export가 그래프를 자르는 지점(Detect 마지막 Conv)과 같다. 예선 CPU 경로와 본선 NPU 경로에서 같은 부분이 양자화된다.",
    ("h", "⑤ 캘리브레이션 방식 비교와 메모리 문제"),
    "MinMax·Percentile·Entropy를 holdout에서 비교해 Percentile을 골랐다(B 95.13 → 96.57, D 90.06 → 90.27). "
    "ORT 히스토그램 보정기는 모든 이미지의 중간 텐서를 메모리에 쌓아 1280×384에서 장당 약 350MB가 늘었다(256장이면 약 90GB). "
    "8장 단위 분할 수집으로 상한을 고정했다. 박스 회귀 Conv까지 FP32로 둔 변형(boxfp32)은 holdout 점수가 같으면서 느려 채택하지 않았다.",
    ("h", "⑥ 같은 INT8 파일의 점수가 CPU마다 달랐다 → 가중치 7비트"),
    "공용 PC(i9-14900KF)에서 만든 INT8을 측정 노트북(i7-10510U)에서 채점하자 FP32는 소수점까지 같은데 INT8만 1~3점 낮았다. 순위도 뒤집혔다(A가 B보다 0.80점 높음 → B가 0.82점 높음). "
    "원인은 uint8×int8 곱을 하는 AVX2 명령 VPMADDUBSW다. 이 명령은 이웃한 두 곱을 int16에 더해 255×127×2 = 64,770이 32,767을 넘으면 잘린다. "
    "VNNI가 있는 x86과 ARM에는 이 문제가 없다. 가중치를 7비트로 줄여(255×64×2 = 32,640) 포화가 생길 수 없게 했다. "
    "재현 검증: VNNI CPU에서 holdout 22개를 다시 채점하니 FP32는 노트북과 같았고, INT8은 최대 0.30점, 주력 B·D는 0.06점 이내였다.",
    ("h", "⑦ P5 헤드 제거"),
    "재학습 없이 학습된 A의 P5 출력만 꺼 보면(근사) holdout mAP가 97.22 → 87.55로 떨어졌다. Car가 98.83 → 78.64로 가장 크게 떨어졌다(가까운 큰 차량). "
    "그래서 1280×384에서는 떼면 깨진다고 예상했지만, P5 경로를 뺀 구조로 다시 학습한 B는 FP32 97.17로 A와 같은 수준이었다. 남은 P4 헤드가 큰 물체를 흡수한 것이다. "
    "640×192에서는 근사에서도 0.54점만 떨어졌다(P5 담당 GT 1.28%).",
    ("h", "⑧ 결정 규칙을 먼저 정하고 eval_val은 마지막에 1회"),
    "규칙: P5 제거 모델의 INT8 mAP가 원본보다 0.5점 이내로 낮으면 느리지 않을 때 채택하고, 0.5점보다 많이 낮으면 10% 이상 빨라야 채택한다(holdout, 4스레드 p50). "
    "B는 A보다 0.19점 낮고 15라운드 측정에서 INT8이 9% 빨라 채택했다. D는 C보다 0.02점 높고 속도가 같거나 빨라 채택했다. "
    "이 결정을 커밋(fcba017)한 뒤 eval_val을 채점했고, 결과를 보고 바꾸지 않았다.",
    ("h", "⑨ 2차 라운드: 더 줄이고 손실을 메우려 했지만 채택하지 않았다"),
    "eval_val에서 INT8 손실이 holdout보다 컸기 때문에(B 1.60, D 1.14점) 두 방향을 더 시도했다. 후보를 만들기 전에 채택 규칙을 DECISION_R2.md(ca8c45c)에 먼저 커밋했다. "
    "규칙은 다음과 같다. 양자화 설정 교체는 holdout INT8 +0.30점 이상이면서 M1 지연 1.05배 이하일 때 채택한다. 프루닝 모델은 손실 0.5점 이내이면서 M1에서 10% 이상 빠를 때 채택한다. "
    "eval_val은 채택된 후보가 있을 때만 다시 1회 채점하기로 했다.",
    ("li", "채널 프루닝(torch-pruning, C2f의 chunk를 Conv 두 개로 펼친 동치 모듈로 바꾼 뒤 L2 크기 기준, 채널 8의 배수) + 같은 설정으로 재학습: "
           "연산량 −26%/−41% 모델 4개 모두 INT8 holdout이 기준보다 0.66~2.79점 낮아 탈락했다. 프루닝은 FP32보다 INT8 손실을 더 키웠다(B: 0.47 → 0.83 → 1.72점)."),
    ("li", "양자화 설정 모델당 10~11가지(박스 경로 FP32, 민감도 상위 층 FP32, Percentile 값, 보정 장수, Entropy, 조합): "
           "보정 이미지 seed만 바꿔도 holdout이 B 0.45점, D 1.64점 움직여 0.30점 기준보다 잡음이 컸다. 그래서 후보와 기준을 seed 3개로 만들어 평균으로 판정했다. "
           "이 판정 방식은 seed 0 결과를 본 뒤에 정한 사후 변경이다. 채택을 더 어렵게 만드는 쪽이며, eval_val은 보지 않았다. "
           "B는 +0.30을 넘는 설정이 없었다. D는 박스 경로 FP32가 평균 +0.83점(세 seed 모두 +0.5 이상)으로 정확도 조건을 통과했다. "
           "그러나 M1 지연이 6.03 → 6.38ms(1.058배, 15라운드 모두 느림)로 속도 조건 1.05배를 넘어 채택하지 않았다. 결과를 보고 기준을 완화하지 않았다."),
])

# ================================================================ 6.2 프로토콜
t = T[12]
set_cell(t.rows[1].cells[1], "1280×384 (정밀도, A→B) / 640×192 (초경량, C→D)")
set_cell(t.rows[2].cells[1], "1")
set_cell(t.rows[3].cells[1], "30회 (측정마다, 결과에서 제외) + 측정 시작 전 120초 예열(지속 성능 상태)")
set_cell(t.rows[4].cells[1], "측정 1회 = 300회 반복. 이를 15라운드 반복하며 라운드마다 모델 순서를 무작위로 섞음(seed 0). 측정마다 새 프로세스")
set_cell(t.rows[5].cells[1], "p50(라운드별 p50의 중앙값)을 지연으로 보고, p99 병기. 평균은 쓰지 않음. FPS = 1000 / p50 (batch 1)\n"
         "두 모델의 속도 비는 라운드 단위 부트스트랩(20,000회) 95% 구간으로 판정")
set_cell(t.rows[6].cells[1], "scripts/bench_suite.py → bench_onnx.py. ONNX Runtime 1.29.0 CPUExecutionProvider, intra_op 4 / inter_op 1 / "
         "ORT_SEQUENTIAL / 그래프 최적화 ALL, 고정 난수 입력(seed 0). 지연은 모델 추론(sess.run)만이며 전처리·후처리 제외. "
         "피크 메모리 = psutil peak RSS. 모델 크기 = ONNX 파일 크기(10^6 바이트 = 1MB). 정확도 = scripts/predict_kitti.py(conf 0.001, NMS IoU 0.65) → kitti_eval.py(공식 devkit 판정 포트)")

# ================================================================ 6.3 정확도
t = T[13]
set_cell(t.rows[0].cells[1], "Baseline A\n(FP32)", bold=True, align=1)
set_cell(t.rows[0].cells[2], "최적화 후 B\n(P5 제거 + INT8)", bold=True, align=1)
for i, k in enumerate(["Car", "Pedestrian", "Cyclist", "mAP"], start=1):
    set_cell(t.rows[i].cells[1], f2(eA[k]), align=1)
    set_cell(t.rows[i].cells[2], f2(eB[k]), align=1)
    set_cell(t.rows[i].cells[3], f"{eA[k] - eB[k]:.2f}", align=1)
caption_before(t, "[정밀도 1280×384] eval_val 1,000장, Moderate AP40")
t2 = clone_after(t._tbl, t, "[초경량 640×192] eval_val 1,000장, Moderate AP40")
set_cell(t2.rows[0].cells[1], "Baseline C\n(FP32)", bold=True, align=1)
set_cell(t2.rows[0].cells[2], "최적화 후 D\n(P5 제거 + INT8)", bold=True, align=1)
for i, k in enumerate(["Car", "Pedestrian", "Cyclist", "mAP"], start=1):
    set_cell(t2.rows[i].cells[1], f2(eC[k]), align=1)
    set_cell(t2.rows[i].cells[2], f2(eD[k]), align=1)
    set_cell(t2.rows[i].cells[3], f"{eC[k] - eD[k]:.2f}", align=1)
# 단계별 분해 표
t3 = clone_after(t2._tbl, t2, "[단계별 분해] 같은 eval_val에서 기법을 하나씩 더했을 때 (Moderate mAP, 괄호는 Car / Ped / Cyc)")
steps = [
    ("원본 FP32 (A / C)", ev(A, "fp32"), ev(C, "fp32")),
    ("+ P5 제거, FP32 (B / D)", eBf, eDf),
    ("+ 전체 INT8", ev(B, "int8"), ev(D, "int8")),
    ("+ Conv만 INT8", ev(B, "int8-conv"), ev(D, "int8-conv")),
    ("+ 혼합 INT8, MinMax", bMin, dMin),
    ("+ 혼합 INT8, Percentile (채택)", eB, eD),
]
fill_grid(t3, [(n, f"{x['mAP']:.2f}  ({x['Car']:.1f} / {x['Pedestrian']:.1f} / {x['Cyclist']:.1f})",
                f"{y['mAP']:.2f}  ({y['Car']:.1f} / {y['Pedestrian']:.1f} / {y['Cyclist']:.1f})", "") for n, x, y in steps],
          header=("단계", "1280×384", "640×192", ""))
# 4열 표에서 마지막 열은 비교 대상: holdout 값
set_cell(t3.rows[0].cells[3], "holdout mAP\n(1280 / 640)", bold=True, align=1)
hv = [(A, C, "fp32"), (B, D, "fp32"), (B, D, "int8"), (B, D, "int8-conv"), (B, D, "int8-mixed"), (B, D, I8)]
for i, (s1, s2, v) in enumerate(hv, start=1):
    set_cell(t3.rows[i].cells[3], f"{hold(s1, v):.2f} / {hold(s2, v):.2f}", align=1)

# ================================================================ 6.4 속도
def speed_table(t, s0, v0, s1, v1, cap, label0, label1):
    set_cell(t.rows[0].cells[1], label0, bold=True, align=1)
    set_cell(t.rows[0].cells[2], label1, bold=True, align=1)
    set_cell(t.rows[1].cells[0], "추론 Latency (ms)\np50 (p99)")
    if HAVE_M1:
        a, b = m1(s0, v0, "p50"), m1(s1, v1, "p50")
        set_cell(t.rows[1].cells[1], f"{a:.2f} ({m1(s0, v0, 'p99'):.2f})", align=1)
        set_cell(t.rows[1].cells[2], f"{b:.2f} ({m1(s1, v1, 'p99'):.2f})", align=1)
        set_cell(t.rows[1].cells[3], f"{(a - b) / a * 100:.1f}", align=1)
        fa, fb = m1(s0, v0, "fps"), m1(s1, v1, "fps")
        set_cell(t.rows[2].cells[1], f"{fa:.1f}", align=1); set_cell(t.rows[2].cells[2], f"{fb:.1f}", align=1)
        set_cell(t.rows[2].cells[3], f"{(fb - fa) / fa * 100:.1f}", align=1)
        ra, rb = m1(s0, v0, "rss"), m1(s1, v1, "rss")
        set_cell(t.rows[4].cells[1], f"{ra:.1f}", align=1); set_cell(t.rows[4].cells[2], f"{rb:.1f}", align=1)
        set_cell(t.rows[4].cells[3], f"{(ra - rb) / ra * 100:.1f}", align=1)
    else:
        for r in (1, 2, 4):
            for c in (1, 2, 3):
                set_cell(t.rows[r].cells[c], "【M1】", align=1)
    za, zb = SIZE[f"{s0}_op13_{v0}"], SIZE[f"{s1}_op13_{v1}"]
    set_cell(t.rows[3].cells[1], f2(za), align=1); set_cell(t.rows[3].cells[2], f2(zb), align=1)
    set_cell(t.rows[3].cells[3], f"{(za - zb) / za * 100:.1f}", align=1)
    caption_before(t, cap)


t = T[14]
speed_table(t, A, "fp32", B, I8, f"[정밀도 1280×384] {m1cpu}, 4스레드", "Baseline A\n(FP32)", "최적화 후 B\n(P5 제거 + INT8)")
t2 = clone_after(t._tbl, t, None)
speed_table(t2, C, "fp32", D, I8, f"[초경량 640×192] {m1cpu}, 4스레드", "Baseline C\n(FP32)", "최적화 후 D\n(P5 제거 + INT8)")
# 보조 측정 PC 표 (i7-10510U, 8개 모델)
t3 = clone_after(t2._tbl, t2, "[보조 측정] Intel Core i7-10510U (AVX2, VNNI 없음), 같은 프로토콜 4스레드 15라운드 — 다른 CPU에서도 같은 경향인지 확인용")
names = [(A, "fp32", "A FP32"), (A, I8, "A INT8"), (B, "fp32", "B FP32"), (B, I8, "B INT8 (주력)"),
         (C, "fp32", "C FP32"), (C, I8, "C INT8"), (D, "fp32", "D FP32"), (D, I8, "D INT8 (주력)")]
rows = []
for s, v, n in names:
    r = LAP[f"{s}_op13_{v}"]
    rows.append((n, f"{r['p50_ms']:.2f} ({r['p99_ms']:.2f})", f"{1000 / r['p50_ms']:.1f}", f"{r['peak_rss_MB']:.1f}"))
fill_grid(t3, rows, header=("모델", "Latency p50 (p99) ms", "FPS", "피크 메모리 MB"))
# 2차 라운드 표 (채택하지 않은 후보)
t4 = clone_after(t3._tbl, t3, "[2차 라운드] 시도했지만 채택하지 않은 개선 — holdout 500장, 기준은 같은 PC에서 다시 양자화한 B 96.70 / D 90.15 (DECISION_R2.md)")
r2rows = [
    ("B 프루닝 m75 (6.95→5.11 GFLOPs)", "96.04 (−0.66)", "측정 안 함", "탈락: 손실 0.5점 초과"),
    ("B 프루닝 m60 (→4.08 GFLOPs)", "94.77 (−1.93)", "측정 안 함", "탈락: 손실 0.5점 초과"),
    ("D 프루닝 m75 (1.74→1.28 GFLOPs)", "88.57 (−1.58)", "측정 안 함", "탈락: 손실 0.5점 초과"),
    ("D 프루닝 m60 (→1.02 GFLOPs)", "87.36 (−2.79)", "측정 안 함", "탈락: 손실 0.5점 초과"),
    ("B 보정 512장 + 박스 경로 FP32 (seed 3회 평균)", "96.98 (+0.27)", "—", "탈락: +0.30 미달"),
    ("D 박스 경로 FP32 (seed 3회 평균)", "90.93 (+0.83)", "6.38 vs 6.03 ms (1.058배)", "탈락: 지연 1.05배 초과"),
]
fill_grid(t4, r2rows, header=("후보", "INT8 holdout mAP (기준 대비)", "M1 INT8 p50", "판정"))

# ================================================================ 6.5 해석
dB = {k: eBf[k] - eB[k] for k in eB}
dD = {k: eDf[k] - eD[k] for k in eD}
lapB = LAP[f"{B}_op13_fp32"]["p50_ms"] / LAP[f"{B}_op13_{I8}"]["p50_ms"]
lapD = LAP[f"{D}_op13_fp32"]["p50_ms"] / LAP[f"{D}_op13_{I8}"]["p50_ms"]
m1ratio = (f"M1 Pro에서는 B {m1(B, 'fp32', 'p50') / m1(B, I8, 'p50'):.2f}배, D {m1(D, 'fp32', 'p50') / m1(D, I8, 'p50'):.2f}배다."
           if HAVE_M1 else "M1 Pro 배율: 【M1】.")
fill_box(T[15], [
    ("h", "정확도 손실은 어디서 생겼나"),
    f"P5 제거 자체는 손실이 아니었다. FP32끼리 비교하면 1280×384에서 A {eA['mAP']:.2f} → B {eBf['mAP']:.2f}({eBf['mAP'] - eA['mAP']:+.2f}), "
    f"640×192에서 C {eC['mAP']:.2f} → D {eDf['mAP']:.2f}({eDf['mAP'] - eC['mAP']:+.2f})로, 학습 시드 하나로는 구분하기 어려운 크기다. "
    f"손실은 거의 전부 INT8 단계에서 나왔다. B는 FP32 대비 {dB['mAP']:.2f}점, D는 {dD['mAP']:.2f}점이다.",
    ("h", "640×192의 Car 손실 — IoU 0.7과 박스 회귀"),
    f"INT8 손실의 클래스별 크기는 B에서 Car {dB['Car']:.2f} / Ped {dB['Pedestrian']:.2f} / Cyc {dB['Cyclist']:.2f}, "
    f"D에서 Car {dD['Car']:.2f} / Ped {dD['Pedestrian']:.2f} / Cyc {dD['Cyclist']:.2f}이다. 1280×384에서는 세 클래스가 비슷하게 떨어졌다. "
    "다만 Cyclist는 eval_val의 Moderate GT가 138개뿐이라 1~2점은 표본 흔들림 안쪽이다. 눈에 띄는 것은 640×192의 Car다. "
    "KITTI는 Car만 IoU 0.7을 요구하고(Ped·Cyc는 0.5), 640×192에서는 물체가 원본의 절반 크기로 줄어 박스 좌표 1~2픽셀 오차가 IoU를 크게 흔든다. "
    "검출은 그대로인데 박스가 조금 어긋나 TP가 FP로 바뀌는 유형이다. 근거는 박스 회귀 Conv까지 FP32로 둔 A의 변형(boxfp32)이다. "
    f"같은 MinMax 보정의 혼합 INT8 대비 Car가 {ev(A, 'int8-mixed')['Car']:.2f} → {ev(A, 'int8-mixed-boxfp32')['Car']:.2f}로 회복됐다. "
    "다만 2차 라운드에서 D의 박스 경로를 FP32로 두었을 때 holdout 이득(+0.83점)은 대부분 Pedestrian에서 나왔다(87.00 → 87.91, seed 3회 평균). Car는 거의 그대로였다. "
    "그래서 640×192 Car 손실을 박스 회귀 양자화 하나로 설명할 수는 없고, 작은 물체의 특징 해상도 자체가 원인일 가능성이 남는다.",
    ("h", "holdout과 eval_val의 차이"),
    f"holdout에서 B의 INT8 손실은 0.60점이었으나 eval_val에서는 {dB['mAP']:.2f}점으로 커졌다. Percentile이 MinMax보다 낫다는 holdout 결론도 eval_val에서는 "
    f"B +{eB['mAP'] - bMin['mAP']:.2f}, D +{eD['mAP'] - dMin['mAP']:.2f}로 좁아졌다. 500장 holdout에서 고른 설정의 이득은 일부가 표본 운이었던 셈이다. "
    f"P5 제거 판단도 마찬가지로, eval_val의 INT8은 B가 A보다 {eAi['mAP'] - eB['mAP']:.2f}점, D가 C보다 {eCi['mAP'] - eD['mAP']:.2f}점 낮다. "
    "사전에 정한 규칙대로 결정은 바꾸지 않고 그대로 보고한다. FP32 점수가 holdout보다 eval_val에서 1.5~2.7점 낮은 것은 네 모델 모두 같으며, 주로 Pedestrian(holdout보다 4~6점 낮음)에서 온다.",
    ("h", "속도 ① INT8 이득은 CPU 명령어와 해상도에 따라 다르다"),
    (f"M1 Pro에서 INT8 배율은 1280×384에서 A {m1r(A,'fp32',A,I8):.2f}배·B {m1r(B,'fp32',B,I8):.2f}배, 640×192에서 C {m1r(C,'fp32',C,I8):.2f}배·D {m1r(D,'fp32',D,I8):.2f}배다. "
     f"같은 파일이 i7-10510U에서는 B {lapB:.2f}배, D {lapD:.2f}배에 그쳤다. 차이는 int8 곱셈 명령에서 온다고 본다. "
     "ARM64(M1)에는 int8 4개를 곱해 32비트에 바로 더하는 내적 명령(SDOT/UDOT)이 있다. VNNI가 없는 AVX2는 16비트 중간합을 거치는 두 단계 명령을 써야 하고, 포화를 피하려고 가중치도 7비트로 줄였다(5.3 ⑥). "
     "두 CPU 모두 640×192에서 배율이 작다. 해상도가 1/4이 되면 Conv 연산은 1/4로 줄지만, FP32로 남긴 디코드 경로·Quantize/Dequantize 변환·스레드 동기화 같은 고정 비용은 그만큼 줄지 않기 때문이다." if HAVE_M1 else
     f"i7-10510U에서 INT8 배율은 B {lapB:.2f}배, D {lapD:.2f}배다. M1 Pro 배율: 【M1】."),
    ("h", "속도 ② P5 제거"),
    (f"P5 제거로 연산량은 1280×384에서 {pct(GFLOPS[A], GFLOPS[B])}, 640×192에서 {pct(GFLOPS[C], GFLOPS[D])} 줄었다. M1에서 B는 A보다 INT8 {m1r(A,I8,B,I8):.2f}배·FP32 {m1r(A,'fp32',B,'fp32'):.2f}배, "
     f"D는 C보다 INT8 {m1r(C,I8,D,I8):.2f}배·FP32 {m1r(C,'fp32',D,'fp32'):.2f}배 빨랐다(15라운드 모두, 95% 구간이 1을 넘음). 연산량 감소와 비슷한 크기의 이득이다. "
     "i7에서는 D와 C의 INT8 차이가 라운드 간 흔들림(10% 안팎)에 묻혀 구분되지 않았다. 흔들림이 2~3%인 M1에서는 분명하게 갈렸다." if HAVE_M1 else "【M1】"),
    ("h", "2차 라운드에서 확인한 것"),
    "첫째, 보정 이미지 seed만 바꿔도 holdout INT8 점수가 B 0.45점, D 1.64점 움직인다. "
    "1차에서 Percentile을 고른 holdout 이득(MinMax 대비 B +1.44, D +0.21점)이 eval_val에서 B +0.08, D +0.14점으로 줄어든 것도 같은 잡음으로 설명된다. "
    "둘째, 박스 경로를 FP32로 남기면 D에서 정확도 +0.83점과 지연 +5.8%를 맞바꾼다. 해당 Conv와 DFL이 정수 연산에서 빠지고 Quantize/Dequantize 경계가 늘기 때문이다. "
    "이 트레이드오프는 본선 Hailo에서 박스 출력층을 16비트로 둘지 정하는 근거가 된다. "
    "셋째, 채널 프루닝은 FP32 손실(B −0.30 ~ −0.68점)보다 INT8 손실을 더 키웠다. 채널이 줄면 층마다 담는 정보가 몰려 양자화 오차에 더 민감해지는 것으로 보인다. "
    "특히 Pedestrian이 가장 크게 떨어졌다.",
    ("h", "크기와 메모리"),
    f"INT8로 파일 크기가 약 1/3이 됐다(B {SIZE[f'{B}_op13_fp32']:.2f} → {SIZE[f'{B}_op13_{I8}']:.2f}MB). P5 제거까지 합치면 A FP32 대비 {pct(SIZE[f'{A}_op13_fp32'], SIZE[f'{B}_op13_{I8}'])}다. "
    "피크 메모리는 ORT 런타임과 1280×384 중간 텐서가 대부분이라 파일 크기만큼 줄지는 않는다.",
])

# ================================================================ 7.1 한계
fill_box(T[16], [
    ("li", "본선용 HEF는 아직 없다. COCO YOLO11n 640×640의 컴파일만 확인했고, B·D(end node 4개, 직사각 입력)의 DFC 컴파일과 Hailo 위 정확도·지연은 미측정이다."),
    ("li", f"INT8 손실이 eval_val에서 holdout보다 크다(B {dB['mAP']:.2f}점, D {dD['mAP']:.2f}점). 특히 640×192의 Car 손실({dD['Car']:.2f}점)은 2차 라운드의 박스 경로 FP32로도 줄지 않아 원인이 남아 있다. QAT는 적용하지 않았다."),
    ("li", "holdout(500장) 판정의 잡음이 크다. 보정 이미지 seed만 바꿔도 INT8 점수가 최대 1.64점 움직여, 0.3점 수준의 설정 차이는 holdout으로 가리기 어렵다. 2차 라운드에서 판정을 seed 3회 평균으로 바꾼 것은 결과를 본 뒤의 사후 변경이다(DECISION_R2.md에 기록)."),
    ("li", "정확도 채점 CPU(x86, VNNI)와 속도 측정 CPU(M1 Pro, ARM)가 다르다. FP32는 CPU와 무관하게 같은 점수가 나오고, INT8은 7비트 가중치로 x86 사이 차이를 0.3점 이내로 줄였다. ARM의 INT8 점수는 직접 확인하지 않았다."),
    ("li", "학습을 모델마다 한 번(시드 1개)만 했다. P5 제거 전후 FP32 차이(±0.3점)는 시드 분산 안쪽일 수 있다."),
    ("li", "보고한 지연은 모델 추론만이다. i7-10510U에서 1280×384 전처리(letterbox·정규화)가 장당 약 26ms로 INT8 추론(약 53ms)의 절반 수준이라, 실제 파이프라인에서는 전처리가 병목이 될 수 있다."),
    ("li", "보조 측정 노트북(i7-10510U)은 발열·전력 한도로 라운드 간 지연이 10% 넘게 흔들리는 칸이 있어 640×192의 D vs C INT8 차이를 구분하지 못했다. 주 측정 M1 Pro는 C FP32 한 칸(10.1%)을 빼면 흔들림이 5% 이내다. 지연은 CPU 한 대(M1 Pro)의 값이며 본선 보드(Raspberry Pi 5) 수치가 아니다."),
])

# ================================================================ 7.2 계획
t = T[17]
plan = [
    ("1주차\n11/6~11/12", "B·D FP32 ONNX → DFC 5.x parse(end node 4개) → optimize(level 2, train 보정 256장) → compile HEF. "
     "profiler로 FPS·지연·리소스 사용률 확인. A·C도 같은 절차로 비교용 HEF 생성",
     "1280×384 리소스 초과·context 증가 → 640×192(D) 우선, 1280은 compression·level 조정"),
    ("2주차\n11/13~11/19", "Raspberry Pi 5 + Hailo-10H 실기기에 HailoRT 설치, HEF 추론 + CPU 후처리(DFL·디코드·NMS) 파이프라인 구성. "
     "eval_val로 HEF 정확도 재채점 → ORT INT8 대비 Drop 비교",
     "DFC 양자화 손실이 ORT보다 큼 → 박스 출력층 16비트(ORT에서 D +0.83점 / 지연 +5.8% 확인), finetune/AdaRound(GPU PC·Colab)"),
    ("3주차\n11/20~11/26", "지연 실측: 전처리·NPU 추론·후처리 분리, p50/p99, 30분 연속 구동 발열 확인. "
     "전처리(리사이즈·정규화)를 NPU normalization 층과 고정 비율 리사이즈로 줄이고, 캡처·추론·후처리 파이프라인 병렬화",
     "CPU 후처리·전처리가 병목 → NMS 상위 k 제한, numpy 벡터화, 스레드 분리"),
    ("마무리\n11/27~12/4", "최종 수치 확정, 재현 패키지(HEF·스크립트·README) 정리, 본선 보고서 작성",
     "일정 지연 → 1·2주차 산출물로 최소 결과 확보"),
]
for i, r in enumerate(plan, start=1):
    for j, v in enumerate(r):
        set_cell(t.rows[i].cells[j], v)

# ================================================================ 8. 체크리스트
t = T[18]
chk = [
    "【팀명】_트랙2_예선결과보고서.pdf",
    "models/best_1280x384_op13_fp32.onnx (A), models/best_640x192_op13_fp32.onnx (C)\n학습 가중치: 【.pt 제출 경로】",
    "models/best-p34_1280x384_op13_int8-mixed-percentile.onnx (B)\nmodels/best-p34_640x192_op13_int8-mixed-percentile.onnx (D)",
    "ONNX opset 13 (ONNX 1.22.0), ONNX Runtime 1.29.0 CPU EP, QDQ INT8. SHA-256: MODELS.sha256",
    "scripts/bench_suite.py · bench_onnx.py · bench_compare.py (속도), predict_kitti.py · kitti_eval.py (정확도), Makefile",
    "requirements.txt",
    "README.md (실행 방법), DECISION.md · DECISION_R2.md (1·2차 결정 기록)",
    "KITTI Object Detection 2012: data_object_image_2.zip, data_object_label_2.zip → data/kitti_raw/training/ (README '데이터 준비')",
]
for i, v in enumerate(chk, start=1):
    set_cell(t.rows[i].cells[2], v)

# ================================================================ 9. 재현
fill_box(T[19], [("code", """
git clone https://github.com/honey0127/mokerton-track2.git
cd mokerton-track2
git checkout 【최종 제출 커밋】
# Windows 는 .venv\\Scripts\\activate
python3.12 -m venv .venv && source .venv/bin/activate
# 전체(학습·내보내기 포함)
pip install -r requirements.txt
# 측정·채점만 재현할 때는 이것으로 충분
pip install onnxruntime==1.29.0 numpy "numba>=0.60" psutil opencv-python-headless pyyaml
# 모델 해시 확인 (macOS: shasum -a 256 -c MODELS.sha256)
sha256sum -c MODELS.sha256
# 평가셋 규격 + AP40 판정 규칙 17종 자체 검증
make selftest
# KITTI 원본을 data/kitti_raw/training/{image_2,label_2} 에 둔 뒤
# split 검증 + data/kitti_yolo/images/{train,holdout,eval_val} 생성
make data KITTI_ROOT=data/kitti_raw
""")])

fill_box(T[20], [
    ("b", "정확도 (eval_val 1,000장, Moderate AP40) ", "— 저장소 루트에서"),
    ("code", """
# Baseline A (FP32, 1280x384)
make eval MODEL=models/best_1280x384_op13_fp32.onnx
# 최적화 B (P5 제거 + INT8, 1280x384)
make eval MODEL=models/best-p34_1280x384_op13_int8-mixed-percentile.onnx
# Baseline C / 최적화 D (640x192)
make eval MODEL=models/best_640x192_op13_fp32.onnx
make eval MODEL=models/best-p34_640x192_op13_int8-mixed-percentile.onnx
# 22개 전체 + 요약표:  make eval-all && make summary
"""),
    ("b", "속도·메모리 (4스레드, 예열 120초, 15라운드, 약 20분)", ""),
    ("code", """
python scripts/bench_suite.py \\
  models/best_1280x384_op13_fp32.onnx \\
  models/best_1280x384_op13_int8-mixed-percentile.onnx \\
  models/best-p34_1280x384_op13_fp32.onnx \\
  models/best-p34_1280x384_op13_int8-mixed-percentile.onnx \\
  models/best_640x192_op13_fp32.onnx \\
  models/best_640x192_op13_int8-mixed-percentile.onnx \\
  models/best-p34_640x192_op13_fp32.onnx \\
  models/best-p34_640x192_op13_int8-mixed-percentile.onnx \\
  --threads 4 --rounds 15 --soak 120 --tag key15 \\
  --jsonl benchmarks/bench_key15_check.jsonl \\
  --raw benchmarks/bench_key15_check_rounds.jsonl
python scripts/bench_compare.py \\
  --jsonl benchmarks/bench_key15_check.jsonl --threads 4
"""),
])

fill_box(T[21], [
    ("b", "make eval: ", "predict_kitti.py가 '[완료] 1000장, 예측 N개, …'와 장당 전처리/추론/후처리 시간을 출력한다. 이어서 kitti_eval.py가 Car/Pedestrian/Cyclist × Easy/Moderate/Hard 표와 "
     "'>>> 공식 점수 (Moderate AP40, 3클래스 평균): XX.XX'를 출력하고 results.jsonl에 기록한다. Moderate 열이 6.3절 표의 값이다"
     f"(A {eA['mAP']:.2f}, B {eB['mAP']:.2f}, C {eC['mAP']:.2f}, D {eD['mAP']:.2f})."),
    ("b", "bench_suite: ", "라운드마다 '[k/120] R… 모델 t=4 p50=… p99=… rss=…MB'를 출력한 뒤 [집계]로 모델별 p50 중앙값과 라운드 간 흔들림(%)을 낸다. "
     "bench_compare는 모델별 p50 중앙값·최소·최대·p99 표와 'INT8 속도배율 A~D', 'B vs A', 'D vs C'의 속도 비와 95% 구간을 출력한다. "
     "6.4절의 Latency는 p50 중앙값(괄호 p99), FPS는 1000/p50, 피크 메모리는 jsonl의 peak_rss_MB, 모델 크기는 size_MB다."),
    ("b", "허용 오차: ", "FP32 정확도는 CPU와 무관하게 소수점 둘째 자리까지 같다. INT8은 CPU 명령어 차이로 ±0.3점 안에서 달라질 수 있다"
     "(우리 측정: AVX2 노트북 vs AVX-512 VNNI에서 holdout 22개 모델 최대 0.30, 주력 0.06). 지연 절대값은 CPU마다 다르므로 속도 비(INT8 배율, B vs A)로 비교한다."),
])

fill_box(T[22], [
    ("li", "모든 명령은 저장소 루트에서 실행한다(data/kitti.yaml의 상대경로가 실행 위치 기준)."),
    ("li", "KITTI 원본 2개 파일(이미지 약 12GB, 라벨 16MB)만 필요하다. 별도로 내려받을 가중치는 없다. 측정·채점은 저장소의 ONNX만 쓴다."),
    ("li", "GPU·CUDA는 필요 없다(CPUExecutionProvider 고정). requirements.txt의 ultralytics가 torch를 함께 설치하므로, 측정·채점만 재현할 때는 9.1의 최소 패키지만 설치해도 된다. 학습·내보내기 재현 시 torch는 PC에 맞는 명령으로 설치한다(requirements.txt 주석)."),
    ("li", "INT8 모델은 가중치 7비트다. VNNI가 없는 x86 CPU에서도 점수가 같게 나오도록 한 설정이며, 다시 양자화할 때 --no-reduce-range를 쓰면 CPU에 따라 점수가 달라진다."),
    ("li", "지연 측정은 전원 연결, 저전력 모드 끔, 다른 앱 종료 상태에서 한다. 예열 120초가 포함되어 15라운드에 약 20분 걸린다. 노트북은 라운드 간 10% 안팎 흔들림이 정상이다."),
    ("li", "macOS에서 onnxruntime 1.29.0은 macOS 14(Sonoma) 이상, Python 3.10 이상에서만 설치된다(PyPI 휠이 macosx_14_0_arm64 전용). macOS 13에서는 pip가 1.23.2까지만 찾는다."),
    ("li", "Apple Silicon에서는 arm64 Python을 써야 한다. Rosetta(x86_64)로 돌리면 지연이 크게 늘어 의미가 없다. "
     "확인: python -c \"import platform; print(platform.machine())\" → arm64"),
    ("li", "eval_val 채점 시간(2코어 Xeon 기준): 1280×384 모델 1개 약 1.5분, 640×192 약 45초. 첫 실행은 numba JIT 컴파일로 수십 초 더 걸린다."),
])

d.save(OUT)
print("saved", OUT, "M1:", HAVE_M1)
